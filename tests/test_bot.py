"""整體流程測試：模擬用戶在 Telegram 裡輸入指令、按按鈕，檢查機器人實際送出的內容。

Telegram 與波場鏈都換成假的，不會真的連網。
"""
import io
import time
import unittest

import httpx
from openpyxl import load_workbook
from telegram import Update
from telegram.error import BadRequest, Conflict, NetworkError, TimedOut
from telegram.ext import CallbackContext

import bot
from bot import STUB_TEXT, build_application, payment_job
from db import DB
from tests.fakes import (
    ALICE,
    BOB,
    OWNER,
    OWNER2,
    SHOP_ADDR,
    TZ,
    FakeRequest,
    callback_update,
    make_address,
    make_cfg,
    message_update,
    tron_item,
)

SCREEN_CALLS = ("editMessageText", "sendMessage", "sendDocument")


def button_texts(message: dict) -> list[str]:
    return [button["text"] for row in message["reply_markup"]["inline_keyboard"] for button in row]


def callback_of(message: dict, label: str) -> str:
    for row in message["reply_markup"]["inline_keyboard"]:
        for button in row:
            if button["text"] == label:
                return button["callback_data"]
    raise AssertionError(f"找不到按鈕「{label}」：{button_texts(message)}")


class BotTestCase(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.req = FakeRequest()
        self.db = DB(":memory:", TZ)
        self.chain: list[dict] = []  # 假的鏈上入帳清單
        self.chain_down = False
        self.chain_requests: list[httpx.Request] = []
        self.http = httpx.AsyncClient(transport=httpx.MockTransport(self._trongrid))
        self.app = build_application(make_cfg(), self.db, request=self.req, http=self.http)
        self.errors: list[BaseException] = []
        self.user_message_id = None

        async def capture(update, context):
            self.errors.append(context.error)

        self.app.add_error_handler(capture)
        await self.app.initialize()

    async def asyncTearDown(self):
        await self.app.shutdown()
        await self.http.aclose()
        self.db.close()

    def _trongrid(self, request: httpx.Request) -> httpx.Response:
        self.chain_requests.append(request)
        if self.chain_down:
            return httpx.Response(503, text="service unavailable")
        return httpx.Response(200, json={"data": self.chain, "success": True, "meta": {"page_size": len(self.chain)}})

    # ---------- 模擬操作 ----------

    async def send(self, user: dict, text: str | None = None, **content) -> None:
        """用戶傳一則訊息。"""
        self.req.calls.clear()
        update = message_update(user, text, **content)
        self.user_message_id = update["message"]["message_id"]
        await self.app.process_update(Update.de_json(update, self.app.bot))
        self.assertEqual(self.errors, [])

    async def click(self, user: dict, data: str, *, message_id: int | None = None) -> None:
        """用戶按按鈕（預設按的是對話中目前那一則訊息）。"""
        self.req.calls.clear()
        update = callback_update(user, data, message_id or self.panel(user) or 500)
        await self.app.process_update(Update.de_json(update, self.app.bot))
        self.assertEqual(self.errors, [])

    async def run_job(self) -> None:
        self.req.calls.clear()
        await payment_job(CallbackContext(self.app))

    async def owner_set(self, field: str, value: str) -> None:
        await self.click(OWNER, f"adm:set:{field}")
        await self.send(OWNER, value)

    # ---------- 檢查工具 ----------

    def names(self) -> list[str]:
        return [name for name, _ in self.req.calls]

    def sent(self, api: str) -> list[dict]:
        return [params for name, params in self.req.calls if name == api]

    def only(self, api: str) -> dict:
        calls = self.sent(api)
        self.assertEqual(len(calls), 1, f"預期剛好一次 {api}，實際：{self.names()}")
        return calls[0]

    def screen(self) -> dict:
        """這次操作之後，對話中那一則訊息的最新內容。"""
        shown = [p for name, p in self.req.calls if name in SCREEN_CALLS and p.get("text") != STUB_TEXT]
        self.assertEqual(len(shown), 1, f"預期畫面只更新一次，實際：{self.names()}")
        return shown[0]

    def panel(self, user: dict) -> int | None:
        """資料庫記錄的「機器人那一則訊息」的編號。"""
        return self.db.get_panel(user["id"])[0]

    def assert_removed(self, user: dict, message_id: int) -> None:
        self.assertIn({"chat_id": user["id"], "message_id": message_id}, self.sent("deleteMessage"))

    def assert_user_message_removed(self, user: dict) -> None:
        self.assert_removed(user, self.user_message_id)

    def count(self, table: str) -> int:
        return self.db.conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]

    def give_points(self, user: dict, points: int) -> None:
        self.db.upsert_user(user["id"], user.get("username"), user["first_name"])
        with self.db.tx():
            self.db._add_points(user["id"], points, "checkin", "測試", int(time.time()))

    def latest_order(self, user: dict):
        return self.db.list_orders(user["id"])[0]

    def status(self, order) -> str:
        return self.db.get_order(order["order_no"])["status"]

    def shift_order(self, order, *, created_ago: int, expires_ago: int) -> int:
        """把訂單的時間往回撥，模擬時間經過。回傳現在時間。"""
        now = int(time.time())
        self.db.conn.execute(
            "UPDATE orders SET created_at = ?, expires_at = ? WHERE id = ?",
            (now - created_ago, now - expires_ago, order["id"]),
        )
        return now


class SingleMessageTest(BotTestCase):
    """對話裡永遠只留機器人的一則訊息。"""

    async def test_first_start_sends_one_message_and_removes_the_command(self):
        await self.send(ALICE, "/start")
        self.assertEqual(sorted(self.names()), ["deleteMessage", "deleteMessages", "sendMessage"])
        menu = self.only("sendMessage")
        self.assert_user_message_removed(ALICE)  # 用戶打的 /start 被清掉
        self.assertEqual(self.panel(ALICE), menu["_message_id"])

    async def test_leftovers_from_before_are_swept_when_no_message_is_on_record(self):
        # 機器人沒有這個對話的記錄（第一次見面，或升級前留下的舊訊息）：發出新的那一則後，把前面的清掉
        await self.send(ALICE, "/start")
        new_id = self.only("sendMessage")["_message_id"]
        sweep = self.only("deleteMessages")
        self.assertEqual(sweep["chat_id"], ALICE["id"])
        self.assertEqual(sweep["message_ids"], list(range(new_id - 100, new_id)))  # 只清它之前的，不含新的那一則

        await self.send(ALICE, "/start")  # 之後已有記錄，就不再做這種清理
        self.assertEqual(self.sent("deleteMessages"), [])

    async def test_sweep_failure_is_harmless(self):
        self.req.errors["deleteMessages"] = "Bad Request: message identifiers are not specified"
        await self.send(ALICE, "/start")
        self.assertEqual(self.only("sendMessage")["text"], "測試中")
        self.assertIsNotNone(self.panel(ALICE))

    async def test_everything_afterwards_edits_that_same_message(self):
        await self.send(ALICE, "/start")
        panel = self.panel(ALICE)
        steps = [
            lambda: self.click(ALICE, "product:A"),
            lambda: self.click(ALICE, "menu"),
            lambda: self.click(ALICE, "orders"),
            lambda: self.send(ALICE, "/start"),
            lambda: self.send(ALICE, "隨便打一句話"),
            lambda: self.send(ALICE, "/unknown"),
            lambda: self.send(ALICE, location={"latitude": 25.0, "longitude": 121.5}),
        ]
        for index, step in enumerate(steps):
            with self.subTest(step=index):
                await step()
                self.assertEqual(self.sent("sendMessage"), [])  # 沒有多出新訊息
                self.assertEqual(self.only("editMessageText")["message_id"], panel)
                self.assertEqual(self.panel(ALICE), panel)

    async def test_every_user_message_is_removed(self):
        await self.send(ALICE, "/start")
        for text in ("/start", "你好", "/unknown"):
            with self.subTest(text=text):
                await self.send(ALICE, text)
                self.assert_user_message_removed(ALICE)
        await self.send(ALICE, location={"latitude": 25.0, "longitude": 121.5})
        self.assert_user_message_removed(ALICE)

    async def test_message_deleted_by_the_user_is_replaced_by_a_new_one(self):
        await self.send(ALICE, "/start")
        old = self.panel(ALICE)
        self.req.errors["editMessageText"] = "Bad Request: message to edit not found"
        self.req.errors["deleteMessage"] = (
            lambda p: "Bad Request: message to delete not found" if p["message_id"] == old else None
        )
        await self.send(ALICE, "/start")
        new = self.only("sendMessage")
        self.assertEqual(new["text"], "測試中")
        self.assertEqual(self.panel(ALICE), new["_message_id"])
        self.assertNotEqual(new["_message_id"], old)

    async def test_pressing_same_button_twice_changes_nothing(self):
        await self.send(ALICE, "/start")
        panel = self.panel(ALICE)
        self.req.errors["editMessageText"] = (
            "Bad Request: message is not modified: specified new message content and reply markup "
            "are exactly the same as a current content and reply markup of the message"
        )
        await self.click(ALICE, "orders")
        self.assertEqual(self.sent("sendMessage"), [])
        self.assertEqual(self.panel(ALICE), panel)

    async def test_clicking_a_leftover_old_message_cleans_it_up(self):
        await self.send(ALICE, "/start")
        panel = self.panel(ALICE)
        await self.click(ALICE, "orders", message_id=77)  # 按到殘留在上面的舊訊息
        self.assert_removed(ALICE, 77)
        self.assertEqual(self.only("editMessageText")["message_id"], panel)  # 畫面更新在目前那一則
        self.assertEqual(self.panel(ALICE), panel)

    async def test_unknown_panel_adopts_the_clicked_message(self):
        # 資料庫沒有記錄（例如舊版留下的訊息）：直接沿用被按的那一則
        await self.click(ALICE, "orders", message_id=77)
        self.assertEqual(self.only("editMessageText")["message_id"], 77)
        self.assertEqual(self.sent("sendMessage"), [])
        self.assertEqual(self.panel(ALICE), 77)

    async def test_notification_arrives_as_a_new_message_and_old_one_is_removed(self):
        await self.send(ALICE, "/start")
        old = self.panel(ALICE)
        self.req.calls.clear()
        await bot._notify(CallbackContext(self.app), ALICE["id"], "通知內容", [bot.BACK])
        new = self.only("sendMessage")  # 新訊息才會讓手機跳通知
        self.assert_removed(ALICE, old)
        self.assertEqual(self.panel(ALICE), new["_message_id"])
        self.assertEqual(self.sent("editMessageText"), [])

    async def test_old_message_too_old_to_delete_becomes_a_one_line_stub(self):
        await self.send(ALICE, "/start")
        old = self.panel(ALICE)
        self.req.errors["deleteMessage"] = (
            lambda p: "Bad Request: message can't be deleted for everyone" if p["message_id"] == old else None
        )
        self.req.calls.clear()
        await bot._notify(CallbackContext(self.app), ALICE["id"], "通知內容", [bot.BACK])
        stub = self.only("editMessageText")
        self.assertEqual((stub["message_id"], stub["text"]), (old, STUB_TEXT))
        self.assertNotIn("reply_markup", stub)  # 舊訊息上的按鈕一併拿掉
        self.assertEqual(self.panel(ALICE), self.only("sendMessage")["_message_id"])

    async def test_notification_failure_is_tolerated(self):
        self.req.errors["sendMessage"] = "Forbidden: bot was blocked by the user"
        await bot._notify(CallbackContext(self.app), BOB["id"], "通知內容", [bot.BACK])
        self.assertIsNone(self.panel(BOB))


class ErrorLogTest(BotTestCase):
    """紀錄檔要分得出「Telegram 那邊暫時連不上」和「機器人自己出錯」。"""

    async def report(self, error: Exception, update=None, job=None) -> None:
        context = CallbackContext.from_error(update, error, self.app, job=job)
        await bot.on_error(update, context)

    async def test_telegram_outage_while_waiting_for_updates_is_a_one_line_warning(self):
        for error in (NetworkError("Bad Gateway"), TimedOut()):
            with self.subTest(error=type(error).__name__):
                with self.assertLogs("bot", level="WARNING") as logs:
                    await self.report(error)
                self.assertEqual(len(logs.records), 1)
                self.assertEqual(logs.records[0].levelname, "WARNING")
                self.assertIn("會自動重試", logs.output[0])
                self.assertIsNone(logs.records[0].exc_info)  # 不附一大段堆疊

    async def test_real_problems_are_still_logged_as_errors_with_details(self):
        update = Update.de_json(message_update(ALICE, "/start"), self.app.bot)
        (payment_job_entry,) = self.app.job_queue.get_jobs_by_name("payment")  # 機器人實際註冊的查帳工作
        cases = [
            ("處理用戶訊息時連線中斷", NetworkError("Bad Gateway"), update, None),
            ("定時工作裡出錯", NetworkError("Bad Gateway"), None, payment_job_entry),
            ("送出的內容被 Telegram 拒絕", BadRequest("can't parse entities"), None, None),
            ("有另一個程式用同一把金鑰在收訊息", Conflict("terminated by other getUpdates request"), None, None),
            ("程式本身的錯誤", ValueError("boom"), None, None),
        ]
        for label, error, upd, job in cases:
            with self.subTest(label):
                with self.assertLogs("bot", level="WARNING") as logs:
                    await self.report(error, update=upd, job=job)
                self.assertEqual(logs.records[0].levelname, "ERROR")
                self.assertIsNotNone(logs.records[0].exc_info)


class MenuTest(BotTestCase):
    async def test_start_shows_message_and_five_buttons_in_required_layout(self):
        await self.send(ALICE, "/start")
        message = self.screen()
        self.assertEqual(message["text"], "測試中")
        rows = message["reply_markup"]["inline_keyboard"]
        self.assertEqual([len(row) for row in rows], [1, 1, 1, 2])  # 按鈕四、五並排
        self.assertEqual(button_texts(message), ["簽到", "商品A", "商品B", "頻道連結", "查詢訂單"])
        self.assertEqual(rows[3][1]["callback_data"], "orders")
        self.assertEqual(self.db.get_user(ALICE["id"])["full_name"], "小明")

    async def test_channel_button_before_link_is_set(self):
        await self.send(ALICE, "/start")
        rows = self.screen()["reply_markup"]["inline_keyboard"]
        self.assertEqual([len(row) for row in rows], [1, 1, 1, 2])
        self.assertNotIn("url", rows[3][0])
        await self.click(ALICE, rows[3][0]["callback_data"])
        self.assertIn("尚未設定", self.only("answerCallbackQuery")["text"])

    async def test_menu_uses_owner_customisations(self):
        self.db.set_setting("welcome_text", "歡迎光臨\n每日簽到領積分")
        self.db.set_setting("channel_url", "https://t.me/example")
        self.db.set_product_name("A", "VIP 月卡")
        self.db.set_product_name("B", "年度方案 <特價>")
        await self.send(ALICE, "/start")
        message = self.screen()
        self.assertEqual(message["text"], "歡迎光臨\n每日簽到領積分")
        self.assertNotIn("parse_mode", message)  # 自訂訊息當純文字顯示，裡面有 < > 也不會出錯
        self.assertEqual(button_texts(message), ["簽到", "VIP 月卡", "年度方案 <特價>", "頻道連結", "查詢訂單"])
        self.assertEqual(message["reply_markup"]["inline_keyboard"][3][0]["url"], "https://t.me/example")

        await self.click(ALICE, "product:B")
        self.assertIn("<b>年度方案 &lt;特價&gt;</b>", self.screen()["text"])  # 名稱裡的符號有正確處理

    async def test_menu_still_shows_if_telegram_rejects_the_channel_link(self):
        self.db.set_setting("channel_url", "https://bad.example/x")

        def reject_bad_url(params):
            return "Bad Request: inline keyboard button URL is invalid" if "bad.example" in str(params) else None

        self.req.errors["sendMessage"] = self.req.errors["editMessageText"] = reject_bad_url
        await self.send(ALICE, "/start")
        self.assertEqual(self.sent("sendMessage")[-1]["text"], "測試中")
        self.assertIsNotNone(self.panel(ALICE))

    async def test_stale_or_unknown_button_returns_to_menu(self):
        await self.send(ALICE, "/start")
        await self.click(ALICE, "something-old")
        self.assertIn("已失效", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.screen()["text"], "測試中")


class CheckinTest(BotTestCase):
    async def test_checkin_once_per_day(self):
        await self.send(ALICE, "/start")
        await self.click(ALICE, "checkin")
        answer = self.only("answerCallbackQuery")
        self.assertIn("簽到成功", answer["text"])
        self.assertIn("目前積分：1", answer["text"])
        self.assertTrue(answer["show_alert"])
        self.assertEqual(self.names(), ["answerCallbackQuery"])  # 用彈出視窗顯示，不產生任何訊息

        await self.click(ALICE, "checkin")
        self.assertIn("已經簽到過", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.get_points(ALICE["id"]), 1)
        self.assertEqual(self.count("checkins"), 1)
        self.assertEqual(self.count("points_ledger"), 1)


class OwnerPageTest(BotTestCase):
    async def test_set_is_recognised_by_owner_id_only(self):
        await self.send(ALICE, "/start")

        await self.send(ALICE, "/whatever")
        unknown_command = (self.names(), self.screen()["text"])
        await self.send(ALICE, "/set")
        not_owner = (self.names(), self.screen()["text"])
        # 不是業主的 ID 輸入 /set：反應和輸入任何看不懂的指令完全一樣
        self.assertEqual(not_owner, unknown_command)
        self.assertEqual(not_owner[1], "測試中")
        self.assertNotIn("業主", str(self.req.calls))

        # 就算有人硬送出業主頁面的按鈕指令，也只會得到和無效按鈕一樣的反應
        for data in ("adm:page", "adm:set:b_price", "adm:set:address", "adm:orders", "adm:export", "adm:addr_ok:abcd1234"):
            with self.subTest(data=data):
                await self.click(ALICE, data)
                self.assertIn("已失效", self.only("answerCallbackQuery")["text"])
                self.assertEqual(self.screen()["text"], "測試中")
                self.assertEqual(self.sent("sendDocument"), [])
                self.assertNotIn("業主", str(self.req.calls))
        await self.send(ALICE, "0.01")  # 一般用戶打數字不會被當成設定值
        self.assertEqual(self.db.get_product("B")["price_units"], 0)

        await self.send(OWNER, "/set")  # 業主的 ID 才會打開頁面
        self.assertIn("業主專用頁面", self.screen()["text"])
        self.assert_user_message_removed(OWNER)

    async def test_set_appears_only_in_owner_command_menu(self):
        await self.send(ALICE, "/start")
        self.assertEqual(self.sent("setMyCommands"), [])

        await self.send(OWNER, "/start")
        scoped = self.only("setMyCommands")
        self.assertEqual(scoped["scope"], {"type": "chat", "chat_id": OWNER["id"]})  # 只套用在業主自己的聊天
        self.assertEqual([c["command"] for c in scoped["commands"]], ["start", "set"])

    async def test_owner_start_still_works_if_command_menu_cannot_be_set(self):
        self.req.errors["setMyCommands"] = "Bad Request: chat not found"
        await self.send(OWNER, "/start")
        self.assertEqual(self.screen()["text"], "測試中")

    async def test_page_layout(self):
        await self.send(OWNER, "/set")
        page = self.screen()
        for expected in ("主選單訊息：測試中", "頻道連結：未設定", "「商品A」", "「商品B」", "未設定（USDT 付款暫不開放）"):
            self.assertIn(expected, page["text"])
        self.assertEqual(
            button_texts(page),
            ["主選單訊息", "頻道連結", "商品A 名稱", "商品B 名稱", "商品A 積分", "商品A 價格",
             "商品B 價格", "簽到積分", "USDT 收款地址", "最近訂單", "匯出表單", "« 返回主選單"],
        )

    async def test_whole_setting_flow_stays_in_one_message(self):
        await self.send(OWNER, "/set")
        panel = self.panel(OWNER)

        await self.click(OWNER, "adm:set:a_points")
        prompt = self.screen()
        self.assertIn("設定「商品A 所需積分」", prompt["text"])
        self.assertIn("目前：未設定", prompt["text"])
        self.assertEqual(button_texts(prompt), ["« 返回業主頁面"])

        await self.send(OWNER, "abc")
        self.assertIn("輸入有誤", self.screen()["text"])
        self.assert_user_message_removed(OWNER)  # 打錯的那則也會清掉

        await self.send(OWNER, "3")
        done = self.screen()
        self.assertIn("✅ 已更新「商品A 所需積分」", done["text"])
        self.assertIn("所需積分：3 積分", done["text"])
        self.assert_user_message_removed(OWNER)
        self.assertEqual(self.db.get_product("A")["points_cost"], 3)

        self.assertEqual(self.panel(OWNER), panel)  # 從頭到尾都是同一則訊息

    async def test_owner_sets_every_field(self):
        await self.send(OWNER, "/set")
        panel = self.panel(OWNER)
        values = {
            "welcome": "歡迎光臨\n每日簽到領積分",
            "channel": "@my_channel",
            "a_name": "VIP 月卡",
            "b_name": "年度方案",
            "a_points": "3",
            "a_price": "9.9",
            "b_price": "19.9",
            "checkin": "2",
        }
        for field, value in values.items():
            with self.subTest(field=field):
                await self.owner_set(field, value)
                self.assertIn(f"✅ 已更新「{bot.FIELDS[field][0]}」", self.screen()["text"])
                self.assertEqual(self.sent("sendMessage"), [])
        self.assertEqual(self.panel(OWNER), panel)

        page = self.screen()
        for expected in ("主選單訊息：歡迎光臨…", "頻道連結：https://t.me/my_channel", "每日簽到積分：2 積分",
                         "「VIP 月卡」", "「年度方案」", "所需積分：3 積分", "USDT 價格：9.90 USDT", "USDT 價格：19.90 USDT"):
            self.assertIn(expected, page["text"])
        self.assertIn("🔗 測試頻道連結", button_texts(page))

        a, b = self.db.get_product("A"), self.db.get_product("B")
        self.assertEqual((a["name"], a["points_cost"], a["price_units"]), ("VIP 月卡", 3, 9_900_000))
        self.assertEqual((b["name"], b["points_cost"], b["price_units"]), ("年度方案", 0, 19_900_000))
        self.assertEqual(self.db.checkin_points(), 2)
        self.assertEqual(self.db.welcome_text(), "歡迎光臨\n每日簽到領積分")
        self.assertEqual(self.db.channel_url(), "https://t.me/my_channel")

        await self.click(OWNER, "adm:set:welcome")  # 設定頁會顯示完整的目前內容
        self.assertIn("歡迎光臨\n每日簽到領積分", self.screen()["text"])

    async def test_invalid_inputs_are_rejected_and_can_be_retried(self):
        bad_inputs = {
            "welcome": ["x" * 1001],
            "channel": ["my_channel", "ftp://example.com", "https://nodot", "@ab", "有 空格"],
            "a_name": ["x" * 31],
            "a_points": ["1.5", "-1", "十", "9999999999"],
            "a_price": ["abc", "-5", "1.234", "1e30"],
            "b_price": ["abc", "9,9"],
            "checkin": ["0", "-1", "abc"],
            "address": ["T123", SHOP_ADDR[:-1] + ("a" if SHOP_ADDR[-1] != "a" else "b"), "0x" + "a" * 40],
        }
        await self.send(OWNER, "/set")
        for field, values in bad_inputs.items():
            await self.click(OWNER, f"adm:set:{field}")
            for value in values:
                with self.subTest(field=field, value=value[:20]):
                    await self.send(OWNER, value)
                    self.assertIn("輸入有誤", self.screen()["text"])
        a, b = self.db.get_product("A"), self.db.get_product("B")
        self.assertEqual((a["name"], a["points_cost"], a["price_units"], b["price_units"]), ("商品A", 0, 0, 0))
        self.assertEqual((self.db.welcome_text(), self.db.channel_url(), self.db.usdt_address()), ("測試中", "", ""))
        self.assertEqual(self.db.checkin_points(), 1)

        await self.click(OWNER, "adm:set:b_price")
        await self.send(OWNER, "abc")
        await self.send(OWNER, "20")  # 輸入有誤之後仍在等輸入，改打正確的值就成功
        self.assertEqual(self.db.get_product("B")["price_units"], 20_000_000)

    async def test_zero_turns_things_off(self):
        self.db.set_points_cost("A", 5)
        self.db.set_price("A", 1_000_000)
        self.db.set_setting("channel_url", "https://t.me/example")
        self.db.set_setting("usdt_address", SHOP_ADDR)
        await self.send(OWNER, "/set")
        for field in ("a_points", "a_price", "channel", "address"):
            await self.owner_set(field, "0")
            self.assertIn("✅ 已更新", self.screen()["text"])
        a = self.db.get_product("A")
        self.assertEqual((a["points_cost"], a["price_units"]), (0, 0))
        self.assertEqual((self.db.channel_url(), self.db.usdt_address()), ("", ""))

    async def test_pending_input_survives_a_bot_restart(self):
        await self.send(OWNER, "/set")
        await self.click(OWNER, "adm:set:b_price")
        # 模擬機器人重啟（例如重新部署）：記憶體裡的狀態全部消失，只剩資料庫
        self.app.user_data[OWNER["id"]].clear()
        self.app.bot_data.pop("locks", None)
        await self.send(OWNER, "19.9")
        self.assertIn("✅ 已更新「商品B 價格」", self.screen()["text"])
        self.assertEqual(self.db.get_product("B")["price_units"], 19_900_000)

        await self.send(OWNER, "5")  # 已經回到業主頁面，再打數字不會被當成設定值
        self.assertEqual(self.db.get_product("B")["price_units"], 19_900_000)

    async def test_back_button_cancels_the_pending_input(self):
        await self.send(OWNER, "/set")
        await self.click(OWNER, "adm:set:a_points")
        await self.click(OWNER, "adm:page")
        self.assertIn("業主專用頁面", self.screen()["text"])
        await self.send(OWNER, "50")  # 已經離開輸入畫面，打數字不會再被當成設定值
        self.assertEqual(self.db.get_product("A")["points_cost"], 0)
        self.assertEqual(self.screen()["text"], "測試中")

    async def test_usdt_address_needs_valid_format_and_explicit_confirmation(self):
        await self.send(OWNER, "/set")
        panel = self.panel(OWNER)
        await self.click(OWNER, "adm:set:address")
        prompt = self.screen()["text"]
        self.assertIn("設定「USDT 收款地址」", prompt)
        self.assertIn("私鑰", prompt)  # 明確提醒不要輸入私鑰

        await self.send(OWNER, f"  {SHOP_ADDR}  ")
        confirm = self.screen()
        self.assert_user_message_removed(OWNER)  # 貼上的地址訊息也會清掉
        self.assertIn("請核對新的收款地址", confirm["text"])
        self.assertIn(f"<code>{SHOP_ADDR}</code>", confirm["text"])
        self.assertEqual(button_texts(confirm), ["確認變更", "取消"])
        self.assertEqual(self.db.usdt_address(), "")  # 還沒按確認之前不會生效

        ok = callback_of(confirm, "確認變更")
        await self.click(OWNER, ok)
        done = self.screen()
        self.assertIn("✅ 已更新「USDT 收款地址」", done["text"])
        self.assertIn(f"<code>{SHOP_ADDR}</code>", done["text"])
        self.assertEqual(self.db.usdt_address(), SHOP_ADDR)
        self.assertEqual(self.panel(OWNER), panel)
        self.assertEqual(self.sent("sendMessage"), [])

        other = make_address("someone-else")
        self.db.set_setting("usdt_address", other)
        await self.click(OWNER, ok)  # 同一個確認按鈕不能重複使用
        self.assertIn("已失效", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.usdt_address(), other)

    async def test_address_confirmation_can_be_cancelled(self):
        await self.send(OWNER, "/set")
        await self.owner_set("address", SHOP_ADDR)
        ok = callback_of(self.screen(), "確認變更")
        await self.click(OWNER, "adm:page")  # 按「取消」
        self.assertIn("業主專用頁面", self.screen()["text"])
        await self.click(OWNER, ok)
        self.assertIn("已失效", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.usdt_address(), "")

    async def test_other_owners_are_alerted_when_address_changes(self):
        self.app.bot_data["cfg"] = make_cfg(owner_ids=frozenset({OWNER["id"], OWNER2["id"]}))
        await self.send(OWNER, "/set")
        await self.owner_set("address", SHOP_ADDR)
        await self.click(OWNER, callback_of(self.screen(), "確認變更"))
        alert = self.only("sendMessage")
        self.assertEqual(alert["chat_id"], OWNER2["id"])
        self.assertIn("USDT 收款地址已變更", alert["text"])
        self.assertIn(SHOP_ADDR, alert["text"])
        self.assertIn(str(OWNER["id"]), alert["text"])  # 看得出是誰改的

        await self.owner_set("address", "0")
        self.assertIn("已清除", self.only("sendMessage")["text"])

    async def test_channel_link_rejected_by_telegram_is_reverted(self):
        self.db.set_setting("channel_url", "https://t.me/good")

        def reject_bad_url(params):
            return "Bad Request: inline keyboard button URL is invalid" if "bad.example" in str(params) else None

        self.req.errors["sendMessage"] = self.req.errors["editMessageText"] = reject_bad_url
        await self.send(OWNER, "/set")
        await self.owner_set("channel", "https://bad.example/x")
        self.assertEqual(self.db.channel_url(), "https://t.me/good")  # 還原成原本的連結
        shown = [p for name, p in self.req.calls if name in SCREEN_CALLS and "bad.example" not in str(p)]
        self.assertIn("Telegram 不接受這個連結", shown[-1]["text"])

        await self.send(OWNER, "@another_channel")  # 仍在等輸入，可以直接重打
        self.assertEqual(self.db.channel_url(), "https://t.me/another_channel")

    async def test_export_is_one_excel_file_that_replaces_the_message(self):
        evil = {"id": 300, "first_name": '=HYPERLINK("http://example.com")', "username": "evil"}
        await self.click(evil, "checkin")
        await self.send(OWNER, "/set")
        page = self.panel(OWNER)

        await self.click(OWNER, "adm:export")
        document = self.only("sendDocument")
        self.assertEqual(self.sent("sendMessage"), [])
        filename, content = document["_files"][0]
        self.assertRegex(filename, r"^表單匯出_\d{8}_\d{4}\.xlsx$")
        self.assertEqual(button_texts(document), ["« 返回業主頁面"])
        self.assert_removed(OWNER, page)  # 檔案取代了原本那一則
        self.assertEqual(self.panel(OWNER), document["_message_id"])

        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book.sheetnames, ["用戶積分", "簽到紀錄", "積分異動", "訂單", "USDT入帳紀錄"])
        sheet = book["簽到紀錄"]
        self.assertEqual([cell.value for cell in sheet[1]], ["編號", "用戶ID", "帳號", "名稱", "簽到日期", "獲得積分", "簽到時間"])
        self.assertEqual(sheet.max_row, 2)
        name_cell = sheet.cell(row=2, column=4)
        self.assertEqual(name_cell.value, evil["first_name"])
        self.assertEqual(name_cell.data_type, "s")  # 存成純文字而不是公式，試算表不會執行它
        self.assertEqual(sheet.cell(row=2, column=2).value, 300)

        # 按返回：檔案訊息沒辦法改成文字，所以移除檔案、換成新的業主頁面
        self.req.errors["editMessageText"] = "Bad Request: there is no text in the message to edit"
        await self.click(OWNER, "adm:page")
        back = self.only("sendMessage")
        self.assertIn("業主專用頁面", back["text"])
        self.assert_removed(OWNER, document["_message_id"])
        self.assertEqual(self.panel(OWNER), back["_message_id"])

    async def test_recent_orders(self):
        await self.send(OWNER, "/set")
        await self.click(OWNER, "adm:orders")
        empty = self.screen()
        self.assertIn("目前沒有已完成的訂單", empty["text"])
        self.assertEqual(button_texts(empty), ["« 返回業主頁面"])

        self.db.set_points_cost("A", 1)
        self.give_points(ALICE, 1)
        order = self.db.redeem_with_points(ALICE["id"], "A")
        await self.click(OWNER, "adm:orders")
        text = self.screen()["text"]
        self.assertIn(order["order_no"], text)
        self.assertIn("1 積分", text)
        self.assertIn(f"tg://user?id={ALICE['id']}", text)  # 可以直接點開買家的對話


class ProductTest(BotTestCase):
    async def test_products_are_closed_until_owner_sets_them(self):
        await self.send(ALICE, "/start")
        for code in ("A", "B"):
            await self.click(ALICE, f"product:{code}")
            view = self.screen()
            self.assertIn("尚未開放", view["text"])
            self.assertEqual(button_texts(view), ["« 返回主選單"])
        await self.click(ALICE, "buy:B")
        self.assertIn("不開放 USDT", self.only("answerCallbackQuery")["text"])
        await self.click(ALICE, "redeem:A")
        self.assertIn("不開放積分兌換", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.count("orders"), 0)

    async def test_usdt_stays_closed_until_receiving_address_is_set(self):
        self.db.set_price("B", 19_900_000)
        await self.send(ALICE, "/start")
        await self.click(ALICE, "product:B")
        self.assertIn("尚未開放", self.screen()["text"])
        await self.click(ALICE, "buy:B")  # 硬送出購買指令也不會開單
        self.assertIn("不開放 USDT", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.count("orders"), 0)

        self.db.set_setting("usdt_address", SHOP_ADDR)
        await self.click(ALICE, "product:B")
        self.assertIn("USDT 付款（19.90 USDT）", button_texts(self.screen()))

    async def test_product_a_offers_points_and_usdt(self):
        self.db.set_points_cost("A", 3)
        self.db.set_price("A", 9_900_000)
        self.db.set_setting("usdt_address", SHOP_ADDR)
        self.give_points(ALICE, 5)
        await self.send(ALICE, "/start")
        await self.click(ALICE, "product:A")
        view = self.screen()
        self.assertIn("積分兌換：3 積分", view["text"])
        self.assertIn("USDT 價格：9.90 USDT", view["text"])
        self.assertIn("你目前的積分：5", view["text"])
        self.assertEqual(button_texts(view), ["用積分兌換（3 積分）", "USDT 付款（9.90 USDT）", "« 返回主選單"])

    async def test_product_b_is_usdt_only(self):
        self.db.set_price("B", 19_900_000)
        self.db.set_setting("usdt_address", SHOP_ADDR)
        self.give_points(ALICE, 100_000)
        await self.send(ALICE, "/start")
        await self.click(ALICE, "product:B")
        view = self.screen()
        self.assertIn("僅支援 TRC-20 USDT", view["text"])
        self.assertEqual(button_texts(view), ["USDT 付款（19.90 USDT）", "« 返回主選單"])

        await self.click(ALICE, "redeem:B")  # 就算有人硬送出兌換指令也不行
        self.assertIn("不開放積分兌換", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.get_points(ALICE["id"]), 100_000)
        self.assertEqual(self.count("orders"), 0)


class RedeemTest(BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.db.set_points_cost("A", 3)
        await self.send(ALICE, "/start")

    async def confirm_button(self, user: dict = ALICE) -> str:
        await self.click(user, "redeem:A")
        confirm = self.screen()
        self.assertIn("確定要用 <b>3</b> 積分兌換「商品A」", confirm["text"])
        return callback_of(confirm, "確認兌換")

    async def test_insufficient_points(self):
        self.give_points(ALICE, 2)
        await self.click(ALICE, "redeem:A")
        self.assertIn("積分不足", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.names(), ["answerCallbackQuery"])
        self.assertEqual(self.count("orders"), 0)

    async def test_redeem_success_and_double_click_is_harmless(self):
        self.give_points(ALICE, 5)
        panel = self.panel(ALICE)
        ok = await self.confirm_button()
        await self.click(ALICE, ok)
        order = self.latest_order(ALICE)
        result = self.only("editMessageText")
        self.assertEqual(result["message_id"], panel)  # 成功畫面也是改在同一則訊息上
        self.assertIn("兌換成功", result["text"])
        self.assertIn(order["order_no"], result["text"])
        self.assertIn("剩餘積分：2", result["text"])
        self.assertEqual((order["status"], order["pay_method"], order["points_spent"]), ("paid", "points", 3))

        notice = self.only("sendMessage")  # 業主收到新訂單通知
        self.assertEqual(notice["chat_id"], OWNER["id"])
        self.assertIn(order["order_no"], notice["text"])
        self.assertIn("積分兌換（3 積分）", notice["text"])
        self.assertIn("今日已完成訂單：1 筆", notice["text"])
        self.assertEqual(button_texts(notice), ["最近訂單", "業主專用頁面", "« 返回主選單"])

        await self.click(ALICE, ok)  # 連點第二下
        self.assertIn("已失效", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.get_points(ALICE["id"]), 2)
        self.assertEqual(self.count("orders"), 1)

    async def test_rejected_when_cost_changed_after_confirm_screen(self):
        self.give_points(ALICE, 10)
        ok = await self.confirm_button()
        self.db.set_points_cost("A", 8)
        await self.click(ALICE, ok)
        self.assertIn("有調整", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.get_points(ALICE["id"]), 10)
        self.assertEqual(self.count("orders"), 0)

    async def test_someone_else_cannot_use_my_confirm_button(self):
        self.give_points(ALICE, 5)
        self.give_points(BOB, 5)
        ok = await self.confirm_button()
        await self.click(BOB, ok)
        self.assertIn("已失效", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.db.get_points(BOB["id"]), 5)

    async def test_owner_buying_own_product_keeps_the_success_screen(self):
        self.give_points(OWNER, 3)
        await self.send(OWNER, "/start")
        ok = await self.confirm_button(OWNER)
        await self.click(OWNER, ok)
        self.assertIn("兌換成功", self.screen()["text"])  # 成功畫面不會被自己的「新訂單通知」蓋掉
        self.assertEqual(self.sent("sendMessage"), [])

    async def test_new_order_notice_replaces_owner_screen_and_cancels_pending_input(self):
        self.give_points(ALICE, 3)
        await self.send(OWNER, "/set")
        await self.click(OWNER, "adm:set:b_price")  # 業主正準備輸入價格
        owner_panel = self.panel(OWNER)

        await self.click(ALICE, await self.confirm_button())
        notice = self.only("sendMessage")
        self.assertEqual(notice["chat_id"], OWNER["id"])
        self.assert_removed(OWNER, owner_panel)  # 業主那邊也維持只有一則
        self.assertEqual(self.panel(OWNER), notice["_message_id"])

        await self.send(OWNER, "20")  # 輸入畫面已被通知取代，之後打的字不會被誤當成價格
        self.assertEqual(self.db.get_product("B")["price_units"], 0)

        await self.click(OWNER, "adm:orders")  # 從通知可以直接看最近訂單
        self.assertIn(self.latest_order(ALICE)["order_no"], self.screen()["text"])


class UsdtTest(BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.db.set_price("B", 19_900_000)
        self.db.set_setting("usdt_address", SHOP_ADDR)

    async def buy(self, user: dict = ALICE):
        if self.panel(user) is None:
            await self.send(user, "/start")
        await self.click(user, "buy:B")
        return self.latest_order(user)

    async def test_payment_instructions(self):
        order = await self.buy()
        pay = self.screen()
        self.assertEqual((order["status"], order["pay_method"], order["amount_units"]), ("pending", "usdt", 19_900_100))
        self.assertEqual(order["pay_address"], SHOP_ADDR)
        self.assertIn(f"<code>{SHOP_ADDR}</code>", pay["text"])
        self.assertIn("<code>19.9001</code> USDT", pay["text"])
        self.assertIn(order["order_no"], pay["text"])
        self.assertIn("TRC-20", pay["text"])
        self.assertEqual(button_texts(pay), ["取消訂單", "« 返回主選單"])

        await self.click(ALICE, "buy:B")  # 再按一次不會重複開單
        self.assertIn("尚未付款的訂單", self.screen()["text"])
        self.assertEqual(self.count("orders"), 1)

        other = await self.buy(BOB)  # 另一位用戶會拿到不同的金額
        self.assertEqual(other["amount_units"], 19_900_200)

    async def test_full_payment_flow(self):
        order = await self.buy()
        pay_panel = self.panel(ALICE)
        await self.run_job()  # 鏈上還沒有入帳
        self.assertEqual(self.req.calls, [])
        self.assertEqual(len(self.chain_requests), 1)
        self.assertIn(f"/accounts/{SHOP_ADDR}/", self.chain_requests[0].url.path)

        now = int(time.time())
        self.chain = [
            tron_item("tx-no-tail", 19_900_000, now),  # 沒帶識別尾數
            tron_item("tx-fake-token", order["amount_units"], now, contract=make_address("fake")),
            tron_item("tx-other-receiver", order["amount_units"], now, to=make_address("other")),
            tron_item("tx-approval", order["amount_units"], now, type_="Approval"),
        ]
        await self.run_job()
        self.assertEqual(self.status(order), "pending")
        self.assertEqual(self.req.calls, [])

        self.chain.append(tron_item("tx-good", order["amount_units"], now))
        await self.run_job()
        paid = self.db.get_order(order["order_no"])
        self.assertEqual((paid["status"], paid["txid"]), ("paid", "tx-good"))
        messages = self.sent("sendMessage")
        self.assertEqual([m["chat_id"] for m in messages], [ALICE["id"], OWNER["id"]])
        self.assertIn("付款成功", messages[0]["text"])
        self.assertIn("19.9001 USDT", messages[0]["text"])
        self.assertIn("tx-good", messages[1]["text"])
        self.assertIn(order["order_no"], messages[1]["text"])
        # 付款成功是新的一則（手機才會跳通知），原本那則付款資訊被移除，對話仍然只有一則
        self.assert_removed(ALICE, pay_panel)
        self.assertEqual(self.panel(ALICE), messages[0]["_message_id"])

        bob_order = await self.buy(BOB)  # 還有別的訂單在等，所以下一輪仍會查鏈
        await self.run_job()
        self.assertEqual(self.req.calls, [])  # 同一筆交易不會重複入帳、重複通知
        self.assertEqual(self.status(bob_order), "pending")
        self.assertEqual(self.count("payments"), 2)  # tx-no-tail（未對應）與 tx-good

    async def test_no_chain_query_when_nothing_is_pending(self):
        await self.run_job()
        self.assertEqual(self.chain_requests, [])

    async def test_watch_status_is_logged_only_when_it_changes(self):
        with self.assertLogs("bot", level="INFO") as logs:
            await self.run_job()
        self.assertIn("暫停查鏈", logs.output[0])  # 啟動後第一輪就會留下紀錄，可確認排程有在跑
        with self.assertNoLogs("bot", level="INFO"):
            await self.run_job()

        order = await self.buy()
        with self.assertLogs("bot", level="INFO") as logs:
            await self.run_job()
        self.assertIn(f"開始監看 1 個收款地址的 USDT 入帳：{SHOP_ADDR}", logs.output[0])
        with self.assertNoLogs("bot", level="INFO"):
            await self.run_job()  # 狀態沒變就不重複寫

        self.chain = [tron_item("tx-paid", order["amount_units"], int(time.time()))]
        await self.run_job()
        with self.assertLogs("bot", level="INFO") as logs:
            await self.run_job()
        self.assertIn("暫停查鏈", logs.output[0])

    async def test_changing_address_does_not_break_orders_already_issued(self):
        old_order = await self.buy(ALICE)
        new_address = make_address("new-wallet")
        self.db.set_setting("usdt_address", new_address)

        new_order = await self.buy(BOB)
        self.assertIn(f"<code>{new_address}</code>", self.screen()["text"])  # 新訂單用新地址
        self.assertEqual((old_order["pay_address"], new_order["pay_address"]), (SHOP_ADDR, new_address))
        await self.click(ALICE, "buy:B")
        self.assertIn(f"<code>{SHOP_ADDR}</code>", self.screen()["text"])  # 舊訂單仍顯示當時給的地址

        now = int(time.time())
        self.chain = [
            tron_item("tx-old", old_order["amount_units"], now, to=SHOP_ADDR),
            tron_item("tx-new", new_order["amount_units"], now, to=new_address),
            tron_item("tx-cross", new_order["amount_units"] + 100, now, to=new_address),
        ]
        self.chain_requests.clear()
        await self.run_job()
        self.assertEqual({r.url.path.split("/")[3] for r in self.chain_requests}, {SHOP_ADDR, new_address})
        self.assertEqual((self.status(old_order), self.status(new_order)), ("paid", "paid"))
        self.assertEqual(self.db.get_order(old_order["order_no"])["txid"], "tx-old")
        self.assertEqual(self.db.get_order(new_order["order_no"])["txid"], "tx-new")

    async def test_right_amount_sent_to_the_other_address_does_not_count(self):
        old_order = await self.buy(ALICE)
        new_address = make_address("new-wallet")
        self.db.set_setting("usdt_address", new_address)
        await self.buy(BOB)
        # 金額是舊訂單的，卻付到新地址：不是當初給客人的地址，不能算
        self.chain = [tron_item("tx-wrong-address", old_order["amount_units"], int(time.time()), to=new_address)]
        await self.run_job()
        self.assertEqual(self.status(old_order), "pending")
        row = self.db.conn.execute("SELECT order_no FROM payments WHERE txid = 'tx-wrong-address'").fetchone()
        self.assertIsNone(row["order_no"])

    async def test_expiry_updates_the_payment_screen_in_place(self):
        order = await self.buy()
        panel = self.panel(ALICE)
        now = self.shift_order(order, created_ago=1860, expires_ago=60)
        await self.run_job()
        self.assertEqual(self.status(order), "expired")
        notice = self.only("editMessageText")
        self.assertEqual(notice["message_id"], panel)
        self.assertIn("已超過付款期限", notice["text"])
        self.assertEqual(self.sent("sendMessage"), [])

        # 用戶其實在期限內付了款，只是鏈上確認比較慢才查得到
        self.chain = [tron_item("tx-in-time", order["amount_units"], now - 90)]
        await self.run_job()
        self.assertEqual(self.status(order), "paid")
        self.assertIn("付款成功", self.sent("sendMessage")[0]["text"])

    async def test_expiry_does_not_interrupt_whatever_else_the_user_is_looking_at(self):
        order = await self.buy()
        await self.click(ALICE, "menu")  # 用戶已經離開付款畫面
        self.shift_order(order, created_ago=1860, expires_ago=60)
        await self.run_job()
        self.assertEqual(self.status(order), "expired")
        self.assertEqual(self.req.calls, [])

    async def test_payment_after_deadline_is_not_matched(self):
        order = await self.buy()
        now = self.shift_order(order, created_ago=1860, expires_ago=60)
        self.chain = [tron_item("tx-too-late", order["amount_units"], now - 30)]
        await self.run_job()
        self.assertEqual(self.status(order), "expired")
        self.assertNotIn("付款成功", str(self.req.calls))
        row = self.db.conn.execute("SELECT order_no FROM payments WHERE txid = 'tx-too-late'").fetchone()
        self.assertIsNone(row["order_no"])  # 有留下入帳紀錄，但沒有對應到訂單，交由業主人工處理

    async def test_restart_catches_up_payment_received_while_offline(self):
        order = await self.buy()
        now = self.shift_order(order, created_ago=7200, expires_ago=5400)
        self.db.conn.execute("UPDATE orders SET status = 'expired' WHERE id = ?", (order["id"],))
        self.chain = [tron_item("tx-offline", order["amount_units"], now - 7000)]
        await self.run_job()  # 重啟後的第一輪會往回補查
        self.assertEqual(self.status(order), "paid")

    async def test_old_expired_order_is_not_watched_once_caught_up(self):
        await self.run_job()  # 第一輪結束後就不再往回補查
        order = await self.buy()
        self.shift_order(order, created_ago=7200, expires_ago=5400)
        self.db.conn.execute("UPDATE orders SET status = 'expired' WHERE id = ?", (order["id"],))
        await self.run_job()
        self.assertEqual(self.chain_requests, [])

    async def test_chain_outage_is_tolerated(self):
        order = await self.buy()
        self.chain_down = True
        await self.run_job()  # 查詢失敗不會讓機器人出錯
        self.assertEqual(self.status(order), "pending")
        self.chain_down = False
        self.chain = [tron_item("tx-after-outage", order["amount_units"], int(time.time()))]
        await self.run_job()
        self.assertEqual(self.status(order), "paid")

    async def test_cancel_order(self):
        order = await self.buy()
        await self.click(BOB, f"cancel:{order['order_no']}")  # 別人不能取消我的訂單
        self.assertIn("無法取消", self.only("answerCallbackQuery")["text"])
        self.assertEqual(self.status(order), "pending")

        await self.click(ALICE, f"cancel:{order['order_no']}")
        self.assertIn("已取消", self.screen()["text"])
        self.assertEqual(self.status(order), "cancelled")
        await self.click(ALICE, f"cancel:{order['order_no']}")
        self.assertIn("無法取消", self.only("answerCallbackQuery")["text"])

        new_order = await self.buy()  # 取消後可以重新下單，而且金額不同
        self.assertNotEqual(new_order["order_no"], order["order_no"])
        self.assertNotEqual(new_order["amount_units"], order["amount_units"])


class OrdersTest(BotTestCase):
    async def test_empty(self):
        await self.send(ALICE, "/start")
        await self.click(ALICE, "orders")
        view = self.screen()
        self.assertIn("目前沒有任何訂單", view["text"])
        self.assertEqual(button_texts(view), ["« 返回主選單"])

    async def test_lists_own_orders_only(self):
        self.db.set_points_cost("A", 3)
        self.db.set_price("B", 19_900_000)
        self.db.set_setting("usdt_address", SHOP_ADDR)
        self.give_points(ALICE, 3)
        await self.send(ALICE, "/start")
        await self.send(BOB, "/start")
        await self.click(ALICE, "redeem:A")
        await self.click(ALICE, callback_of(self.screen(), "確認兌換"))
        await self.click(ALICE, "buy:B")
        await self.click(BOB, "buy:B")

        await self.click(ALICE, "orders")
        text = self.screen()["text"]
        mine = self.db.list_orders(ALICE["id"])
        self.assertEqual(len(mine), 2)
        for order in mine:
            self.assertIn(order["order_no"], text)
        self.assertIn("已完成", text)
        self.assertIn("3 積分", text)
        self.assertIn("待付款", text)
        self.assertIn("19.9001 USDT", text)
        self.assertNotIn(self.latest_order(BOB)["order_no"], text)


if __name__ == "__main__":
    unittest.main()
