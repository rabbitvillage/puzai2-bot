"""Telegram 機器人主程式：主選單、簽到、商品、訂單查詢、業主專用頁面、USDT 入帳偵測。

畫面原則：對話裡永遠只留「機器人的一則訊息」。
- 用戶按按鈕或傳訊息：直接修改那一則訊息；用戶自己傳的訊息看完就刪掉。
- 機器人主動通知（付款成功、新訂單）：發一則新的再刪掉舊的，手機才會跳通知。

啟動方式：python bot.py
"""
from __future__ import annotations

import asyncio
import functools
import html
import io
import logging
import re
import secrets
import time
from datetime import datetime

import httpx
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from telegram import BotCommand, BotCommandScopeChat, InlineKeyboardMarkup, LinkPreviewOptions, Update
from telegram import InlineKeyboardButton as Btn
from telegram.constants import ParseMode
from telegram.error import BadRequest, NetworkError, TelegramError
from telegram.ext import (
    Application,
    ApplicationBuilder,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

import config
import tron
from db import (
    DB,
    EXPORTS,
    MAX_POINTS,
    RESERVE_GRACE,
    STARTUP_GRACE,
    STATUS_TEXT,
    InsufficientPoints,
    NoAmountSlot,
    NotAvailable,
    PriceChanged,
    fmt_usdt,
    parse_price,
)

log = logging.getLogger("bot")
HTML = ParseMode.HTML
Ctx = ContextTypes.DEFAULT_TYPE
NO_PREVIEW = LinkPreviewOptions(is_disabled=True)
BACK = [Btn("« 返回主選單", callback_data="menu")]
OWNER_BACK = [Btn("« 返回業主頁面", callback_data="adm:page")]
PUBLIC_COMMANDS = [BotCommand("start", "開啟選單")]
OWNER_COMMANDS = PUBLIC_COMMANDS + [BotCommand("set", "業主專用頁面")]
STUB_TEXT = "（此訊息已更新，請看下方最新訊息）"
SWEEP_SPAN = 100  # Telegram 一次最多能批次刪除 100 則


# ============================== 共用小工具 ==============================


def _cfg(context: Ctx) -> config.Config:
    return context.bot_data["cfg"]


def _db(context: Ctx) -> DB:
    return context.bot_data["db"]


def _touch(update: Update, context: Ctx):
    """每次互動都更新用戶資料，確保用戶表裡有這個人。"""
    user = update.effective_user
    _db(context).upsert_user(user.id, user.username, user.full_name)
    return user


def _is_owner(update: Update, context: Ctx) -> bool:
    return update.effective_user.id in _cfg(context).owner_ids


# ============================== 單一訊息畫面 ==============================


def _lock(context: Ctx, chat_id: int) -> asyncio.Lock:
    """每個對話一把鎖：避免「用戶按按鈕」和「定時通知」同時在換那一則訊息。"""
    return context.bot_data.setdefault("locks", {}).setdefault(chat_id, asyncio.Lock())


async def _retire(context: Ctx, chat_id: int, message_id: int) -> None:
    """讓一則舊訊息退場：能刪就刪；超過 48 小時 Telegram 不給刪，就改成一行提示並拿掉按鈕。"""
    try:
        await context.bot.delete_message(chat_id, message_id)
        return
    except TelegramError as e:
        if "not found" in str(e).lower():  # 本來就不在了
            return
    try:
        await context.bot.edit_message_text(STUB_TEXT, chat_id=chat_id, message_id=message_id)
    except TelegramError:
        try:  # 檔案訊息不能改文字，至少把按鈕拿掉
            await context.bot.edit_message_reply_markup(chat_id=chat_id, message_id=message_id, reply_markup=None)
        except TelegramError:
            pass


async def _sweep(context: Ctx, chat_id: int, below_id: int) -> None:
    """機器人不確定這個對話裡還有沒有舊訊息時（第一次見面，或舊版留下的），把前面最近 100 則清掉。

    只會刪到「這個對話」裡的訊息；編號不屬於這個對話的會被 Telegram 自動略過。
    """
    ids = [i for i in range(below_id - SWEEP_SPAN, below_id) if i > 0]
    try:
        await context.bot.delete_messages(chat_id, ids)
    except TelegramError as e:
        log.info("清理舊訊息時被略過：%s", e)


async def _edit(context: Ctx, chat_id: int, message_id: int, text: str, markup, parse_mode) -> bool:
    """修改既有訊息。成功（或內容本來就一樣）回傳 True；改不了回傳 False。"""
    try:
        await context.bot.edit_message_text(
            text,
            chat_id=chat_id,
            message_id=message_id,
            reply_markup=markup,
            parse_mode=parse_mode,
            link_preview_options=NO_PREVIEW,
        )
        return True
    except BadRequest as e:
        if "not modified" in str(e).lower():
            return True
        log.info("無法修改訊息 %s（%s），改發新訊息", message_id, e)
        return False


async def render(
    context: Ctx,
    chat_id: int,
    text: str,
    rows,
    *,
    tag: str = "",
    parse_mode: str | None = HTML,
    fresh: bool = False,
    document: tuple[bytes, str] | None = None,
    clicked: int | None = None,
) -> None:
    """把對話中「機器人唯一的那一則訊息」換成新內容。所有畫面都從這裡出去。

    - 一般情況：直接修改原訊息，對話不會多出新訊息。
    - fresh=True（主動通知）或 document（檔案）：發一則新的，再讓舊的退場。
    - 原訊息改不了（被用戶刪掉、原本是檔案訊息）：同上。
    - clicked 是用戶按到的那則訊息；如果它不是目前記錄的那一則，代表是殘留的舊訊息，順手清掉。
    - 完全沒有記錄（第一次見面、或舊版留下的對話）：發新的之後，把前面殘留的訊息一併清掉。
    """
    db = _db(context)
    markup = rows if isinstance(rows, InlineKeyboardMarkup) else InlineKeyboardMarkup(rows)
    async with _lock(context, chat_id):
        stored = db.get_panel(chat_id)[0]
        if clicked and stored and clicked != stored:
            await _retire(context, chat_id, clicked)
        target = stored or clicked
        if target and not fresh and document is None:
            if await _edit(context, chat_id, target, text, markup, parse_mode):
                db.set_panel(chat_id, target, tag)
                return
        if document is not None:
            data, filename = document
            message = await context.bot.send_document(
                chat_id, document=data, filename=filename, caption=text, reply_markup=markup, parse_mode=parse_mode
            )
        else:
            message = await context.bot.send_message(
                chat_id, text, reply_markup=markup, parse_mode=parse_mode, link_preview_options=NO_PREVIEW
            )
        db.set_panel(chat_id, message.message_id, tag)
        if target:
            await _retire(context, chat_id, target)
        else:
            await _sweep(context, chat_id, message.message_id)


async def _screen(update: Update, context: Ctx, text: str, rows, *, tag: str = "", parse_mode: str | None = HTML):
    """回應用戶的操作：把那一則訊息換成新畫面。"""
    query = update.callback_query
    clicked = query.message.message_id if query and query.message else None
    await render(context, update.effective_chat.id, text, rows, tag=tag, parse_mode=parse_mode, clicked=clicked)


async def _consume(update: Update, context: Ctx) -> None:
    """刪掉用戶剛傳來的那則訊息，對話才不會一直往上滾。"""
    try:
        await context.bot.delete_message(update.effective_chat.id, update.effective_message.message_id)
    except TelegramError as e:
        if "not found" not in str(e).lower():  # 已經被前面的清理一併刪掉就不用理會
            log.info("無法刪除用戶訊息：%s", e)


async def _notify(context: Ctx, chat_id: int, text: str, rows) -> None:
    """主動通知：發一則新訊息取代原本那一則。對方封鎖機器人等狀況只記錄，不中斷流程。"""
    try:
        await render(context, chat_id, text, rows, fresh=True)
    except TelegramError as e:
        log.warning("無法通知 %s：%s", chat_id, e)


# ============================== 主選單 ==============================


def main_menu(context: Ctx, with_url: bool = True) -> InlineKeyboardMarkup:
    cfg, db = _cfg(context), _db(context)
    url = db.channel_url() if with_url else ""
    if url:
        channel = Btn(cfg.channel_button_text, url=url)
    else:
        channel = Btn(cfg.channel_button_text, callback_data="nochannel")
    return InlineKeyboardMarkup(
        [
            [Btn("簽到", callback_data="checkin")],
            [Btn(db.get_product("A")["name"], callback_data="product:A")],
            [Btn(db.get_product("B")["name"], callback_data="product:B")],
            [channel, Btn("查詢訂單", callback_data="orders")],  # 按鈕四、五並排
        ]
    )


async def show_menu(update: Update, context: Ctx) -> None:
    text = _db(context).welcome_text()
    try:
        await _screen(update, context, text, main_menu(context), tag="menu", parse_mode=None)
    except BadRequest as e:
        # 保險：萬一頻道連結被 Telegram 拒絕，主選單仍然要能顯示
        log.error("主選單顯示失敗（%s），改用不含頻道連結的版本", e)
        await _screen(update, context, text, main_menu(context, with_url=False), tag="menu", parse_mode=None)


async def cmd_start(update: Update, context: Ctx) -> None:
    _touch(update, context)
    context.user_data.clear()
    await show_menu(update, context)
    await _consume(update, context)
    if _is_owner(update, context):
        # 只讓業主自己的指令選單多出 /set；其他用戶的選單裡看不到這個指令
        try:
            await context.bot.set_my_commands(OWNER_COMMANDS, scope=BotCommandScopeChat(update.effective_chat.id))
        except TelegramError as e:
            log.warning("無法設定業主的指令選單：%s", e)


async def on_menu(update: Update, context: Ctx) -> None:
    _touch(update, context)
    context.user_data.clear()
    await update.callback_query.answer()
    await show_menu(update, context)


async def on_nochannel(update: Update, context: Ctx) -> None:
    await update.callback_query.answer("頻道連結尚未設定。", show_alert=True)


async def on_unknown_callback(update: Update, context: Ctx) -> None:
    _touch(update, context)
    await update.callback_query.answer("這個按鈕已失效，已為你回到主選單。", show_alert=True)
    await show_menu(update, context)


async def on_message(update: Update, context: Ctx) -> None:
    """用戶傳來的任何訊息：業主正在設定時當作輸入值，其餘一律回到主選單；看完就刪掉那則訊息。"""
    user = _touch(update, context)
    text = update.effective_message.text or ""
    field = _awaiting_field(context, user.id)
    if field and _is_owner(update, context) and text and not text.startswith("/"):
        await _owner_input(update, context, field, text)
    else:
        await show_menu(update, context)
    await _consume(update, context)


def _awaiting_field(context: Ctx, user_id: int) -> str | None:
    """業主的畫面如果正停在某個設定項目的輸入說明，回傳那個項目。

    以「畫面上現在顯示什麼」為準（記在資料庫），所以機器人重啟也不會忘記正在等輸入；
    畫面一換掉（返回、被通知取代、回主選單），就自然不再等輸入。
    """
    tag = _db(context).get_panel(user_id)[1]
    field = tag.removeprefix("set:") if tag.startswith("set:") else ""
    return field if field in FIELDS else None


# ============================== 簽到 ==============================


async def on_checkin(update: Update, context: Ctx) -> None:
    user = _touch(update, context)
    ok, gained, balance = _db(context).checkin(user.id)
    if ok:
        text = f"簽到成功！獲得 {gained} 積分\n目前積分：{balance}"
    else:
        text = f"今天已經簽到過了，明天再來吧！\n目前積分：{balance}"
    await update.callback_query.answer(text, show_alert=True)


# ============================== 商品 ==============================


async def on_product(update: Update, context: Ctx) -> None:
    user = _touch(update, context)
    await update.callback_query.answer()
    db = _db(context)
    code = context.match.group(1)
    product = db.get_product(code)
    cost = product["points_cost"]
    price = product["price_units"] if db.usdt_address() else 0  # 還沒有收款地址就不顯示 USDT 付款

    lines = [f"<b>{html.escape(product['name'])}</b>", ""]
    rows = []
    if cost > 0:
        lines.append(f"積分兌換：{cost} 積分")
        rows.append([Btn(f"用積分兌換（{cost} 積分）", callback_data=f"redeem:{code}")])
    if price > 0:
        lines.append(f"USDT 價格：{fmt_usdt(price)} USDT")
        rows.append([Btn(f"USDT 付款（{fmt_usdt(price)} USDT）", callback_data=f"buy:{code}")])
    if cost > 0:
        lines += ["", f"你目前的積分：{db.get_points(user.id)}"]
    elif price > 0:
        lines.append("付款方式：僅支援 TRC-20 USDT")
    else:
        lines.append("此商品尚未開放，請稍後再來。")
    rows.append(BACK)
    await _screen(update, context, "\n".join(lines), rows)


async def on_redeem(update: Update, context: Ctx) -> None:
    """積分兌換第一步：顯示確認畫面。"""
    query = update.callback_query
    user = _touch(update, context)
    db = _db(context)
    code = context.match.group(1)
    product = db.get_product(code)
    cost, points = product["points_cost"], db.get_points(user.id)
    if cost <= 0:
        await query.answer("此商品目前不開放積分兌換。", show_alert=True)
        return
    if points < cost:
        await query.answer(f"積分不足：需要 {cost} 積分，你目前有 {points} 積分。", show_alert=True)
        return
    await query.answer()
    token = secrets.token_hex(4)
    context.user_data["redeem"] = (token, code, cost)
    await _screen(
        update,
        context,
        f"確定要用 <b>{cost}</b> 積分兌換「{html.escape(product['name'])}」嗎？\n\n"
        f"目前積分：{points}\n兌換後剩餘：{points - cost}",
        [
            [Btn("確認兌換", callback_data=f"redeem_ok:{token}")],
            [Btn("取消", callback_data=f"product:{code}")],
        ],
    )


async def on_redeem_ok(update: Update, context: Ctx) -> None:
    """積分兌換第二步：真正扣積分、建立訂單。"""
    query = update.callback_query
    user = _touch(update, context)
    db = _db(context)
    # 確認碼只能用一次：連點兩下不會被扣兩次積分
    token, code, cost = context.user_data.pop("redeem", None) or (None, None, None)
    if token is None or token != context.match.group(1):
        await query.answer("這個確認已失效，請重新操作。", show_alert=True)
        return
    try:
        order = db.redeem_with_points(user.id, code, expected_cost=cost)
    except InsufficientPoints:
        await query.answer("積分不足，無法兌換。", show_alert=True)
        return
    except NotAvailable:
        await query.answer("此商品目前不開放積分兌換。", show_alert=True)
        return
    except PriceChanged:
        await query.answer("所需積分剛剛有調整，請重新操作。", show_alert=True)
        return
    await query.answer()
    await _screen(
        update,
        context,
        "\n".join(
            [
                "✅ <b>兌換成功</b>",
                "",
                f"訂單編號：<code>{order['order_no']}</code>",
                f"商品：{html.escape(db.get_product(code)['name'])}",
                f"扣除積分：{order['points_spent']}",
                f"剩餘積分：{db.get_points(user.id)}",
                "",
                "已通知商家處理你的訂單。",
            ]
        ),
        [BACK],
    )
    await _notify_owners(context, order)


def payment_text(context: Ctx, order, reused: bool) -> str:
    db = _db(context)
    lines = ["<b>請完成付款</b>"]
    if reused:
        lines.append("（你已有一筆尚未付款的訂單，以下是那一筆的付款資訊）")
    lines += [
        "",
        f"訂單編號：<code>{order['order_no']}</code>",
        f"商品：{html.escape(db.get_product(order['product_code'])['name'])}",
        "網路：TRON（TRC-20）",
        "收款地址：",
        f"<code>{order['pay_address']}</code>",
        f"應付金額：<code>{fmt_usdt(order['amount_units'], 4)}</code> USDT",
        f"付款期限：{db.fmt_time(order['expires_at'])}",
        "",
        "<b>注意事項</b>",
        "1. 到帳金額必須與「應付金額」完全相同（含小數），系統才能自動對帳。",
        "2. 只接受 TRC-20 網路的 USDT，轉錯網路無法找回。",
        "3. 從交易所提幣會被扣手續費，請確認「實際到帳」等於應付金額。",
        "4. 付款後約 1–2 分鐘會自動通知，不需要回傳截圖。",
    ]
    return "\n".join(lines)


async def on_buy(update: Update, context: Ctx) -> None:
    """USDT 付款：建立待付款訂單並顯示付款資訊。"""
    query = update.callback_query
    user = _touch(update, context)
    cfg, db = _cfg(context), _db(context)
    code = context.match.group(1)
    try:
        order, created = db.create_usdt_order(user.id, code, db.usdt_address(), cfg.order_expire_minutes * 60)
    except NotAvailable:
        await query.answer("此商品目前不開放 USDT 付款。", show_alert=True)
        return
    except NoAmountSlot:
        await query.answer("目前下單人數較多，請稍後再試。", show_alert=True)
        return
    await query.answer()
    await _screen(
        update,
        context,
        payment_text(context, order, reused=not created),
        [[Btn("取消訂單", callback_data=f"cancel:{order['order_no']}")], BACK],
        tag=f"pay:{order['order_no']}",
    )


async def on_cancel_order(update: Update, context: Ctx) -> None:
    query = update.callback_query
    user = _touch(update, context)
    order_no = context.match.group(1)
    if not _db(context).cancel_order(order_no, user.id):
        await query.answer("這筆訂單已無法取消（可能已付款、已逾期或已取消）。", show_alert=True)
        return
    await query.answer()
    await _screen(update, context, f"訂單 <code>{order_no}</code> 已取消，請勿再付款。", [BACK])


# ============================== 查詢訂單 ==============================


def _paid_with(order) -> str:
    if order["pay_method"] == "points":
        return f"{order['points_spent']} 積分"
    return f"{fmt_usdt(order['amount_units'], 4)} USDT"


async def on_orders(update: Update, context: Ctx) -> None:
    user = _touch(update, context)
    await update.callback_query.answer()
    db = _db(context)
    orders = db.list_orders(user.id)
    lines = ["<b>我的訂單</b>" + ("（最近 10 筆）" if orders else "")]
    if not orders:
        lines += ["", "目前沒有任何訂單。"]
    for order in orders:
        status = STATUS_TEXT[order["status"]]
        if order["status"] == "pending":
            status += f"（請於 {db.fmt_time(order['expires_at'], '%H:%M')} 前付款）"
        lines += [
            "",
            f"<code>{order['order_no']}</code>｜{status}",
            f"{html.escape(order['product_name'])}｜{_paid_with(order)}｜{db.fmt_time(order['created_at'])}",
        ]
    await _screen(update, context, "\n".join(lines), [BACK])


# ============================== 業主專用頁面（/set）==============================

# 可設定的項目：代碼 -> （名稱, 輸入說明）
FIELDS = {
    "welcome": ("主選單訊息", "請輸入主選單上方要顯示的訊息（最多 1000 字，可以換行）。"),
    "channel": ("頻道連結", "請輸入頻道連結，例如 https://t.me/頻道名稱 或 @頻道名稱。\n輸入 0 表示清除連結。"),
    "a_name": ("商品A 名稱", "請輸入商品A 要顯示的名稱（最多 30 字）。"),
    "b_name": ("商品B 名稱", "請輸入商品B 要顯示的名稱（最多 30 字）。"),
    "a_points": ("商品A 所需積分", "請輸入商品A 兌換所需的積分（整數）。\n輸入 0 表示關閉積分兌換。"),
    "a_price": ("商品A 價格", "請輸入商品A 的 USDT 價格（最多 2 位小數，例如 9.9）。\n輸入 0 表示關閉 USDT 付款。"),
    "b_price": ("商品B 價格", "請輸入商品B 的 USDT 價格（最多 2 位小數，例如 19.9）。\n輸入 0 表示暫停販售。"),
    "checkin": ("每日簽到積分", "請輸入每次簽到可獲得的積分（1 以上的整數）。"),
    "address": (
        "USDT 收款地址",
        "請貼上收款用的 TRC-20 地址（T 開頭、共 34 碼）。\n輸入 0 表示關閉 USDT 付款。\n\n"
        "機器人只會查詢這個地址的入帳紀錄，不需要、也請不要輸入私鑰或助記詞。",
    ),
}
MAX_WELCOME = 1000
MAX_NAME = 30


def owner_only(handler):
    """業主頁面的按鈕只認業主的 ID；其他人按到，反應和按到無效按鈕一模一樣，不透露這是業主功能。"""

    @functools.wraps(handler)
    async def wrapper(update: Update, context: Ctx) -> None:
        if not _is_owner(update, context):
            await on_unknown_callback(update, context)
            return
        await handler(update, context)

    return wrapper


def _points_label(value: int) -> str:
    return f"{value} 積分" if value > 0 else "未設定（不開放）"


def _price_label(units: int) -> str:
    return f"{fmt_usdt(units)} USDT" if units > 0 else "未設定（不開放）"


def _preview(text: str, limit: int = 30) -> str:
    """多行或太長的文字只顯示開頭，完整內容到設定頁再看。"""
    first, _, rest = text.partition("\n")
    return first[:limit] + ("…" if rest or len(first) > limit else "")


def owner_menu(context: Ctx) -> InlineKeyboardMarkup:
    rows = [
        [Btn("主選單訊息", callback_data="adm:set:welcome"), Btn("頻道連結", callback_data="adm:set:channel")],
        [Btn("商品A 名稱", callback_data="adm:set:a_name"), Btn("商品B 名稱", callback_data="adm:set:b_name")],
        [Btn("商品A 積分", callback_data="adm:set:a_points"), Btn("商品A 價格", callback_data="adm:set:a_price")],
        [Btn("商品B 價格", callback_data="adm:set:b_price"), Btn("簽到積分", callback_data="adm:set:checkin")],
        [Btn("USDT 收款地址", callback_data="adm:set:address")],
        [Btn("最近訂單", callback_data="adm:orders"), Btn("匯出表單", callback_data="adm:export")],
    ]
    channel = _db(context).channel_url()
    if channel:  # 讓業主可以直接點開確認連結沒填錯
        rows.append([Btn("🔗 測試頻道連結", url=channel)])
    rows.append(BACK)
    return InlineKeyboardMarkup(rows)


def owner_text(context: Ctx) -> str:
    db = _db(context)
    a, b, stats = db.get_product("A"), db.get_product("B"), db.stats()
    address, channel = db.usdt_address(), db.channel_url()
    return "\n".join(
        [
            "<b>業主專用頁面</b>",
            "",
            f"主選單訊息：{html.escape(_preview(db.welcome_text()))}",
            f"頻道連結：{html.escape(channel) if channel else '未設定'}",
            f"每日簽到積分：{db.checkin_points()} 積分",
            "",
            f"<b>商品A</b>「{html.escape(a['name'])}」",
            f"　所需積分：{_points_label(a['points_cost'])}",
            f"　USDT 價格：{_price_label(a['price_units'])}",
            f"<b>商品B</b>「{html.escape(b['name'])}」",
            f"　USDT 價格：{_price_label(b['price_units'])}",
            "",
            "USDT 收款地址：",
            f"<code>{address}</code>" if address else "未設定（USDT 付款暫不開放）",
            "",
            f"用戶數：{stats['users']}　今日簽到：{stats['checkins_today']}",
            f"已完成訂單：{stats['orders_paid']}　待付款訂單：{stats['orders_pending']}",
        ]
    )


def _current_value(db: DB, field: str) -> str:
    """設定項目目前的值（已處理成可以安全放進訊息的文字）。"""
    a, b = db.get_product("A"), db.get_product("B")
    if field == "welcome":
        return "\n" + html.escape(db.welcome_text())
    if field == "channel":
        return html.escape(db.channel_url()) or "未設定"
    if field == "a_name":
        return html.escape(a["name"])
    if field == "b_name":
        return html.escape(b["name"])
    if field == "a_points":
        return _points_label(a["points_cost"])
    if field == "a_price":
        return _price_label(a["price_units"])
    if field == "b_price":
        return _price_label(b["price_units"])
    if field == "checkin":
        return f"{db.checkin_points()} 積分"
    address = db.usdt_address()
    return f"\n<code>{address}</code>" if address else "未設定"


def prompt_text(context: Ctx, field: str) -> str:
    label, hint = FIELDS[field]
    return (
        f"<b>設定「{label}」</b>\n\n"
        f"目前：{_current_value(_db(context), field)}\n\n"
        f"{hint}\n\n"
        "直接輸入並送出即可（你送出的訊息會自動清除）。"
    )


def _parse_int(text: str, minimum: int) -> int:
    if not text.isdigit():
        raise ValueError("請輸入整數")
    value = int(text)
    if value < minimum:
        raise ValueError(f"不可小於 {minimum}")
    if value > MAX_POINTS:
        raise ValueError(f"不可超過 {MAX_POINTS}")
    return value


def parse_channel_url(text: str) -> str:
    """接受 @頻道名稱 或完整網址，回傳可放在按鈕上的網址。"""
    if re.fullmatch(r"@[A-Za-z0-9_]{4,32}", text):
        return f"https://t.me/{text[1:]}"
    if re.fullmatch(r"https?://[^\s/]+\.[^\s/]+(/\S*)?", text) and len(text) <= 200:
        return text
    raise ValueError("請輸入 https://t.me/頻道名稱 或 @頻道名稱")


def apply_setting(db: DB, field: str, text: str) -> None:
    """把業主輸入的內容存成設定；格式不對會丟 ValueError（訊息直接回給業主看）。

    收款地址不走這裡（它要先經過核對畫面），只有「輸入 0 關閉」會進來。
    """
    text = text.strip()
    if field == "welcome":
        if not text:
            raise ValueError("訊息不可為空白")
        if len(text) > MAX_WELCOME:
            raise ValueError(f"訊息最多 {MAX_WELCOME} 字，目前 {len(text)} 字")
        db.set_setting("welcome_text", text)
    elif field == "channel":
        db.set_setting("channel_url", "" if text == "0" else parse_channel_url(text))
    elif field in ("a_name", "b_name"):
        if not text or "\n" in text:
            raise ValueError("名稱不可為空白，也不能換行")
        if len(text) > MAX_NAME:
            raise ValueError(f"名稱最多 {MAX_NAME} 字，目前 {len(text)} 字")
        db.set_product_name(field[0].upper(), text)
    elif field == "a_points":
        db.set_points_cost("A", _parse_int(text, minimum=0))
    elif field == "a_price":
        db.set_price("A", parse_price(text))
    elif field == "b_price":
        db.set_price("B", parse_price(text))
    elif field == "checkin":
        db.set_setting("checkin_points", str(_parse_int(text, minimum=1)))
    elif field == "address" and text == "0":
        db.set_setting("usdt_address", "")
    else:
        raise ValueError("未知的設定項目")


async def show_owner_page(update: Update, context: Ctx, notice: str = "") -> None:
    text = (notice + "\n\n" if notice else "") + owner_text(context)
    await _screen(update, context, text, owner_menu(context), tag="owner")


async def cmd_set(update: Update, context: Ctx) -> None:
    """/set：業主專用頁面。只認業主名單裡的 Telegram ID。

    不是業主的人輸入 /set，機器人的反應和收到任何看不懂的訊息完全一樣，不透露有這個頁面。
    """
    if not _is_owner(update, context):
        await on_message(update, context)
        return
    _touch(update, context)
    context.user_data.clear()
    await show_owner_page(update, context)
    await _consume(update, context)


@owner_only
async def on_owner_page(update: Update, context: Ctx) -> None:
    context.user_data.clear()
    await update.callback_query.answer()
    await show_owner_page(update, context)


@owner_only
async def on_owner_field(update: Update, context: Ctx) -> None:
    """業主按下某個設定項目：畫面換成輸入說明，接下來他傳的那則訊息就是新的值。"""
    field = context.match.group(1)
    if field not in FIELDS:
        await on_unknown_callback(update, context)
        return
    await update.callback_query.answer()
    await _screen(update, context, prompt_text(context, field), [OWNER_BACK], tag=f"set:{field}")


async def _owner_input(update: Update, context: Ctx, field: str, text: str) -> None:
    """處理業主輸入的設定值。"""
    db = _db(context)
    label = FIELDS[field][0]
    previous_channel = db.channel_url()
    try:
        if field == "address" and text.strip() != "0":
            await _ask_address_confirmation(update, context, text.strip())
            return
        apply_setting(db, field, text)
        await show_owner_page(update, context, f"✅ 已更新「{label}」。")
    except ValueError as e:
        await _screen(
            update,
            context,
            f"⚠️ 輸入有誤：{html.escape(str(e))}\n\n{prompt_text(context, field)}",
            [OWNER_BACK],
            tag=f"set:{field}",  # 仍停在輸入畫面，可以直接重打
        )
        return
    except BadRequest as e:
        if field != "channel":
            raise
        # 業主頁面上有「測試頻道連結」按鈕；Telegram 不接受這個網址時會在這裡失敗，把設定還原
        log.warning("Telegram 不接受頻道連結：%s", e)
        db.set_setting("channel_url", previous_channel)
        await _screen(
            update,
            context,
            f"⚠️ Telegram 不接受這個連結，請確認網址是否正確。\n\n{prompt_text(context, field)}",
            [OWNER_BACK],
            tag=f"set:{field}",
        )
        return
    log.info("業主 %s 更新了「%s」", update.effective_user.id, label)
    if field == "address":
        await _announce_address_change(update, context, "")


async def _ask_address_confirmation(update: Update, context: Ctx, address: str) -> None:
    """收款地址是最重要的設定：先檢查格式，再讓業主親眼核對一次才生效。"""
    if not tron.is_valid_address(address):
        raise ValueError("這不是有效的 TRC-20 地址。請確認是 T 開頭、共 34 碼，而且沒有多貼或漏貼字元")
    token = secrets.token_hex(4)
    context.user_data["pending_address"] = (token, address)
    await _screen(
        update,
        context,
        "<b>請核對新的收款地址</b>\n\n"
        f"<code>{address}</code>\n\n"
        "請逐字比對開頭與結尾幾碼，確認和你錢包裡的地址完全相同。\n\n"
        "確認後，新建立的訂單會改用這個地址收款；已經建立、還在等付款的訂單仍使用原本的地址。",
        [[Btn("確認變更", callback_data=f"adm:addr_ok:{token}")], [Btn("取消", callback_data="adm:page")]],
    )


@owner_only
async def on_owner_address_ok(update: Update, context: Ctx) -> None:
    query = update.callback_query
    token, address = context.user_data.pop("pending_address", None) or (None, None)
    if token is None or token != context.match.group(1):
        await query.answer("這個確認已失效，請重新設定。", show_alert=True)
        await show_owner_page(update, context)
        return
    _db(context).set_setting("usdt_address", address)
    log.info("業主 %s 把 USDT 收款地址改為 %s", update.effective_user.id, address)
    await query.answer()
    await show_owner_page(update, context, "✅ 已更新「USDT 收款地址」。")
    await _announce_address_change(update, context, address)


async def _announce_address_change(update: Update, context: Ctx, address: str) -> None:
    """收款地址被改動時通知其他業主：萬一有人的帳號被盜用，其他人能馬上發現。"""
    actor = update.effective_user
    text = "\n".join(
        [
            "⚠️ <b>USDT 收款地址已變更</b>",
            "",
            f"操作者：{html.escape(actor.full_name)}（ID <code>{actor.id}</code>）",
            "新的收款地址：",
            f"<code>{address}</code>" if address else "（已清除，USDT 付款關閉）",
            "",
            "如果這不是預期中的變更，請立刻到業主專用頁面檢查。",
        ]
    )
    for owner_id in _cfg(context).owner_ids:
        if owner_id != actor.id:
            await _notify(context, owner_id, text, [[Btn("業主專用頁面", callback_data="adm:page")], BACK])


@owner_only
async def on_owner_orders(update: Update, context: Ctx) -> None:
    await update.callback_query.answer()
    db = _db(context)
    orders = db.recent_paid_orders()
    lines = ["<b>最近完成的訂單</b>" + ("（最多 10 筆）" if orders else "")]
    if not orders:
        lines += ["", "目前沒有已完成的訂單。"]
    for order in orders:
        buyer = html.escape(order["full_name"] or str(order["user_id"]))
        lines += [
            "",
            f"<code>{order['order_no']}</code>｜{html.escape(order['product_name'])}",
            f"{_paid_with(order)}｜<a href=\"tg://user?id={order['user_id']}\">{buyer}</a>"
            f"｜{db.fmt_time(order['paid_at'], '%m-%d %H:%M')}",
        ]
    await _screen(update, context, "\n".join(lines), [OWNER_BACK])


def _width(text: str) -> int:
    """估算文字在試算表裡的寬度（中文字算兩格）。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def xlsx_bytes(db: DB) -> bytes:
    """把所有表單放進同一個 Excel 檔，每張表單一個工作表。"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in EXPORTS:
        title, headers, rows = db.export(name)
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                # 用戶名稱等文字是外部輸入：強制當成純文字，就算開頭是「=」也不會被試算表當公式執行
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            longest = max([_width(header)] + [_width(str(row[index - 1] or "")) for row in rows])
            sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 60)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


@owner_only
async def on_owner_export(update: Update, context: Ctx) -> None:
    """匯出表單：一個 Excel 檔，直接取代對話中的那一則訊息。"""
    query = update.callback_query
    await query.answer("匯出中…")
    db = _db(context)
    titles = "、".join(EXPORTS[name][0] for name in EXPORTS)
    await render(
        context,
        update.effective_chat.id,
        f"表單已匯出，共 {len(EXPORTS)} 張工作表：{titles}。\n\n"
        "請先把檔案存下來；按「返回業主頁面」後，這個檔案會從對話中移除。",
        [OWNER_BACK],
        tag="export",
        parse_mode=None,
        document=(xlsx_bytes(db), f"表單匯出_{db.fmt_time(int(time.time()), '%Y%m%d_%H%M')}.xlsx"),
        clicked=query.message.message_id,
    )


# ============================== USDT 入帳偵測與通知 ==============================


async def _notify_owners(context: Ctx, order) -> None:
    """訂單完成（積分兌換成功或 USDT 入帳）時通知業主，方便安排出貨。"""
    cfg, db = _cfg(context), _db(context)
    user = db.get_user(order["user_id"])
    name = html.escape(user["full_name"] or str(user["user_id"]))
    account = f"（@{html.escape(user['username'])}）" if user["username"] else ""
    lines = [
        "🔔 <b>新訂單（已完成）</b>",
        "",
        f"訂單編號：<code>{order['order_no']}</code>",
        f"用戶：<a href=\"tg://user?id={user['user_id']}\">{name}</a>{account}",
        f"用戶 ID：<code>{user['user_id']}</code>",
        f"商品：{html.escape(db.get_product(order['product_code'])['name'])}",
    ]
    if order["pay_method"] == "points":
        lines.append(f"付款方式：積分兌換（{order['points_spent']} 積分）")
    else:
        lines += [
            "付款方式：USDT（TRC-20）",
            f"實收金額：{fmt_usdt(order['amount_units'], 4)} USDT",
            f"鏈上交易序號：<code>{html.escape(order['txid'])}</code>",
        ]
    lines += ["", f"今日已完成訂單：{db.stats()['orders_paid_today']} 筆（可按「最近訂單」查看全部）"]
    rows = [[Btn("最近訂單", callback_data="adm:orders"), Btn("業主專用頁面", callback_data="adm:page")], BACK]
    for owner_id in cfg.owner_ids:
        if owner_id == order["user_id"]:
            continue  # 業主自己下的單：他已經看到成功畫面，不用再蓋掉它
        # 通知會取代業主當下的畫面；如果他原本停在輸入畫面，畫面換掉後就不再等輸入，
        # 之後打的字不會被誤當成設定值
        await _notify(context, owner_id, "\n".join(lines), rows)


async def payment_job(context: Ctx) -> None:
    """定時工作：先查鏈上入帳並對應訂單，再把過期的訂單標成逾期。

    順序是刻意的：先對帳再處理逾期，剛好壓線付款的人才不會先被標成逾期。
    """
    db = _db(context)
    now = int(time.time())
    await _check_incoming(context, now)
    for order in db.expire_orders(now):
        # 只有當用戶的畫面還停在這張訂單的付款資訊時才更新它，避免打斷他正在做的其他事
        if db.get_panel(order["user_id"])[1] != f"pay:{order['order_no']}":
            continue
        try:
            await render(
                context,
                order["user_id"],
                f"⌛ 訂單 <code>{order['order_no']}</code> 已超過付款期限，請勿再付款。\n\n"
                "若你已在期限內完成付款，入帳確認後仍會自動通知你。",
                [BACK],
            )
        except TelegramError as e:
            log.warning("無法更新 %s 的逾期畫面：%s", order["user_id"], e)


async def _check_incoming(context: Ctx, now: int) -> None:
    cfg, db = _cfg(context), _db(context)
    # 重啟後的第一輪往回多查一點，把停機期間進來的款項補上
    grace = RESERVE_GRACE if context.bot_data.get("caught_up") else STARTUP_GRACE
    targets = db.watch_targets(now, grace)  # 沒有訂單在等付款就不查鏈，節省查詢額度
    # 只在「監看對象改變」時寫一行紀錄，方便確認機器人有沒有在看帳，又不會每 15 秒洗版
    watching = [address for address, _ in targets]
    if watching != context.bot_data.get("watching"):
        context.bot_data["watching"] = watching
        if watching:
            log.info("開始監看 %d 個收款地址的 USDT 入帳：%s", len(watching), "、".join(watching))
        else:
            log.info("目前沒有等待付款的訂單，暫停查鏈。")
    all_ok = True
    for address, since in targets:
        try:
            transfers = await tron.fetch_incoming(
                context.bot_data["http"],
                api_url=cfg.tron_api_url,
                address=address,
                contract=cfg.usdt_contract,
                since_ts=since,
                api_key=cfg.trongrid_api_key,
            )
        except (httpx.HTTPError, tron.TronError, ValueError) as e:
            log.warning("查詢鏈上入帳失敗，下一輪再試：%s", e)
            all_ok = False
            continue
        for transfer in transfers:
            order = db.record_transfer(
                transfer.txid, transfer.from_addr, transfer.to_addr, transfer.amount_units, transfer.block_ts
            )
            if order:
                log.info("訂單 %s 已收到 USDT 付款（交易 %s）", order["order_no"], transfer.txid)
                await _on_usdt_paid(context, order)
    if all_ok:
        context.bot_data["caught_up"] = True


async def _on_usdt_paid(context: Ctx, order) -> None:
    db = _db(context)
    await _notify(
        context,
        order["user_id"],
        "\n".join(
            [
                "✅ <b>付款成功</b>",
                "",
                f"訂單編號：<code>{order['order_no']}</code>",
                f"商品：{html.escape(db.get_product(order['product_code'])['name'])}",
                f"金額：{fmt_usdt(order['amount_units'], 4)} USDT",
                "",
                "已通知商家處理你的訂單。",
            ]
        ),
        [BACK],
    )
    await _notify_owners(context, order)


# ============================== 組裝與啟動 ==============================


async def on_error(update: object, context: Ctx) -> None:
    error = context.error
    transient = isinstance(error, NetworkError) and not isinstance(error, BadRequest)
    if update is None and context.job is None and transient:
        # 「等 Telegram 送訊息過來」的長連線偶爾會被對方伺服器中斷（例如 Bad Gateway、逾時），
        # 套件會自動重連，不是機器人壞掉，所以只留一行警告
        log.warning("與 Telegram 的連線暫時異常，會自動重試：%s", error)
        return
    log.error("處理訊息或定時工作時發生錯誤", exc_info=error)


async def _post_init(app: Application) -> None:
    await app.bot.set_my_commands(PUBLIC_COMMANDS)


async def _post_shutdown(app: Application) -> None:
    await app.bot_data["http"].aclose()
    app.bot_data["db"].close()


def build_application(
    cfg: config.Config, db: DB, *, request=None, http: httpx.AsyncClient | None = None
) -> Application:
    builder = ApplicationBuilder().token(cfg.bot_token).post_init(_post_init).post_shutdown(_post_shutdown)
    if request is not None:  # 測試時換成假的傳輸層，不會真的連到 Telegram
        builder = builder.request(request)
    app = builder.build()
    app.bot_data.update(cfg=cfg, db=db, http=http or httpx.AsyncClient(timeout=15))

    private = filters.ChatType.PRIVATE
    app.add_handler(CommandHandler("start", cmd_start, filters=private))
    app.add_handler(CommandHandler("set", cmd_set, filters=private))
    app.add_handler(CallbackQueryHandler(on_menu, pattern=r"^menu$"))
    app.add_handler(CallbackQueryHandler(on_checkin, pattern=r"^checkin$"))
    app.add_handler(CallbackQueryHandler(on_product, pattern=r"^product:(A|B)$"))
    app.add_handler(CallbackQueryHandler(on_redeem, pattern=r"^redeem:(A|B)$"))
    app.add_handler(CallbackQueryHandler(on_redeem_ok, pattern=r"^redeem_ok:([0-9a-f]+)$"))
    app.add_handler(CallbackQueryHandler(on_buy, pattern=r"^buy:(A|B)$"))
    app.add_handler(CallbackQueryHandler(on_cancel_order, pattern=r"^cancel:([A-Z0-9]+)$"))
    app.add_handler(CallbackQueryHandler(on_orders, pattern=r"^orders$"))
    app.add_handler(CallbackQueryHandler(on_nochannel, pattern=r"^nochannel$"))
    app.add_handler(CallbackQueryHandler(on_owner_page, pattern=r"^adm:page$"))
    app.add_handler(CallbackQueryHandler(on_owner_field, pattern=r"^adm:set:(\w+)$"))
    app.add_handler(CallbackQueryHandler(on_owner_address_ok, pattern=r"^adm:addr_ok:([0-9a-f]+)$"))
    app.add_handler(CallbackQueryHandler(on_owner_orders, pattern=r"^adm:orders$"))
    app.add_handler(CallbackQueryHandler(on_owner_export, pattern=r"^adm:export$"))
    app.add_handler(CallbackQueryHandler(on_unknown_callback))
    app.add_handler(MessageHandler(private, on_message))  # 其餘任何訊息（含看不懂的指令）
    app.add_error_handler(on_error)

    app.job_queue.run_repeating(payment_job, interval=cfg.poll_interval_seconds, first=5, name="payment")
    return app


def main() -> None:
    logging.basicConfig(format="%(asctime)s %(levelname)s %(name)s: %(message)s", level=logging.INFO)
    # 這兩個套件的一般訊息會把含金鑰的網址、或每 15 秒一次的排程紀錄印出來，只留警告以上
    logging.getLogger("httpx").setLevel(logging.WARNING)
    logging.getLogger("apscheduler").setLevel(logging.WARNING)
    try:
        cfg = config.load()
    except config.ConfigError as e:
        raise SystemExit(f"設定錯誤：{e}") from None
    # 紀錄檔的時間改用設定的時區顯示（主機本身是 UTC），和訂單、查詢結果的時間才對得起來
    logging.Formatter.converter = lambda *args: datetime.fromtimestamp(args[-1], cfg.tz).timetuple()
    if not cfg.owner_ids:
        log.warning("尚未設定 OWNER_IDS，目前沒有任何人能使用業主專用頁面 /set。")
    db = DB(cfg.db_path, cfg.tz)
    if not db.usdt_address():
        log.warning("尚未設定 USDT 收款地址，USDT 付款暫不開放（業主可在 /set 頁面設定）。")
    app = build_application(cfg, db)
    log.info("機器人啟動中，按 Ctrl+C 可停止。")
    app.run_polling(allowed_updates=["message", "callback_query"])


if __name__ == "__main__":
    main()
